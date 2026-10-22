"""Text extraction for documents and web pages."""
from __future__ import annotations

import io
import re
from pathlib import Path

DOC_EXT = {".pdf", ".docx", ".pptx", ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json",
           ".log", ".html", ".htm", ".xml", ".yaml", ".yml", ".ini", ".cfg", ".toml", ".srt", ".vtt",
           ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".kt", ".cs", ".cpp", ".c", ".h", ".rs",
           ".go", ".rb", ".php", ".lua", ".sh", ".ps1", ".bat", ".sql", ".css"}


def read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        return raw.decode("utf-16", errors="replace")
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def extract_pdf(path: Path) -> tuple[list[str], dict]:
    """Returns per-page text and metadata. Pages with no text layer come back empty."""
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    pages = []
    for p in reader.pages:
        try:
            pages.append(p.extract_text() or "")
        except Exception:
            pages.append("")
    meta = {"pages": len(pages)}
    try:
        info = reader.metadata or {}
        if info.get("/Title"):
            meta["doc_title"] = str(info["/Title"])
        if info.get("/Author"):
            meta["author"] = str(info["/Author"])
    except Exception:
        pass
    return pages, meta


def render_pdf_page(path: Path, index: int, dpi: int = 150) -> bytes | None:
    """Rasterise one PDF page to JPEG (for OCR via the vision model)."""
    try:
        import pypdfium2 as pdfium
    except Exception:
        return None
    pdf = pdfium.PdfDocument(str(path))
    try:
        page = pdf[index]
        img = page.render(scale=dpi / 72).to_pil().convert("RGB")
        img.thumbnail((1600, 1600))
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=88)
        return buf.getvalue()
    finally:
        pdf.close()


def extract_docx(path: Path) -> str:
    import docx
    d = docx.Document(str(path))
    parts: list[str] = []
    for block in d.element.body.iterchildren():
        tag = block.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = "".join(t.text or "" for t in block.iter() if t.tag.endswith("}t"))
            if text.strip():
                parts.append(text)
        elif tag == "tbl":
            for row in block.iter():
                if row.tag.endswith("}tr"):
                    cells = []
                    for cell in row.iter():
                        if cell.tag.endswith("}tc"):
                            cells.append(" ".join(t.text or "" for t in cell.iter() if t.tag.endswith("}t")).strip())
                    if any(cells):
                        parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def extract_pptx(path: Path) -> list[str]:
    from pptx import Presentation
    prs = Presentation(str(path))
    slides = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    texts.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(f"Speaker notes: {notes}")
        slides.append("\n".join(texts))
    return slides


def html_to_text(html: str, url: str | None = None) -> tuple[str, str | None]:
    """Main-content extraction. trafilatura first (boilerplate removal), BeautifulSoup fallback."""
    title = None
    try:
        import trafilatura
        text = trafilatura.extract(html, url=url, include_tables=True, include_comments=False,
                                   favor_recall=True, output_format="txt")
        try:
            md = trafilatura.extract_metadata(html, default_url=url)
            title = md.title if md else None
        except Exception:
            pass
        if text and len(text) > 200:
            return text, title
    except Exception:
        pass
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html, "html.parser")
    if soup.title and soup.title.string:
        title = title or soup.title.string.strip()
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "aside", "form", "svg"]):
        tag.decompose()
    text = soup.get_text("\n")
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip(), title


def extract_file(path: Path) -> tuple[list[str], dict]:
    """Return a list of 'pages' (one entry for flat documents) plus metadata."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf(path)
    if ext == ".docx":
        return [extract_docx(path)], {}
    if ext == ".pptx":
        slides = extract_pptx(path)
        return slides, {"pages": len(slides), "page_label": "slide"}
    if ext in (".html", ".htm"):
        text, title = html_to_text(read_text_file(path))
        return [text], ({"doc_title": title} if title else {})
    return [read_text_file(path)], {}
